import re
from typing import List, Tuple

from pptx import Presentation


def normalize_text(text: str) -> str:
    text = text.replace("\x0b", "\n")
    text = re.sub(r"\n\s*\n+", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def get_shape_position(shape) -> Tuple[int, int]:
    """获取shape的位置（left, top），用于从左到右、从上到下排序"""
    try:
        # 改为 (left, top) 优先从左到右排序
        return (shape.left, shape.top)
    except AttributeError:
        return (0, 0)


def extract_text_from_shape(shape) -> str:
    """递归提取shape中的所有文字，包括表格、分组对象等"""
    fragments = []

    # 处理普通文本框
    if getattr(shape, "has_text_frame", False):
        text = normalize_text(shape.text or "")
        if text:
            fragments.append(text)

    # 处理表格
    if getattr(shape, "has_table", False):
        table = shape.table
        for row in table.rows:
            row_texts = []
            for cell in row.cells:
                text = normalize_text(cell.text or "")
                if text:
                    row_texts.append(text)
            if row_texts:
                fragments.append(" ".join(row_texts))

    # 处理分组对象（递归，并按位置排序）
    if getattr(shape, "shapes", None):
        sub_shapes = sorted(shape.shapes, key=get_shape_position)
        for sub_shape in sub_shapes:
            sub_text = extract_text_from_shape(sub_shape)
            if sub_text:
                fragments.append(sub_text)

    return "\n".join(fragments).strip()


def extract_slide_texts(pptx_path: str) -> List[str]:
    presentation = Presentation(pptx_path)
    slide_texts: List[str] = []

    for slide in presentation.slides:
        # 按照从左到右、从上到下的顺序排序shapes
        sorted_shapes = sorted(slide.shapes, key=get_shape_position)

        fragments: List[str] = []
        for shape in sorted_shapes:
            text = extract_text_from_shape(shape)
            if text:
                fragments.append(text)

        slide_texts.append("\n".join(fragments).strip())

    return slide_texts


def get_slide_text(pptx_path: str, slide_index: int) -> str:
    slide_texts = extract_slide_texts(pptx_path)
    if 1 <= slide_index <= len(slide_texts):
        return slide_texts[slide_index - 1]
    return ""
